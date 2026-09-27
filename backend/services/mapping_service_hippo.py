"""
Modulo di mappatura dei chunks.
Genera output Excel con tutte le informazioni richieste.
"""

import os
import json
import re
from typing import List, Dict, Any, Tuple, Optional
from sentence_transformers import SentenceTransformer
from dotenv import load_dotenv
import sys

# Try importing from the file we just created
try:
    from hippo import HippoRAGLite
    print("✅ hippo.py importato con successo.")
except ImportError:
    print("⚠️ Please create hippo.py in the same directory!")

# Add parent directory to path to import modules
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config.prompts_config import MAPPING_SYSTEM_PROMPT, get_mapping_user_prompt
from llm_provider import llm_chat

# Load environment variables
load_dotenv()
import numpy as np
from sklearn.metrics.pairwise import cosine_similarity
from sklearn.feature_extraction.text import TfidfVectorizer
from pathlib import Path
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

class ImprovedRequirementsMapper:
    # [FIX 1] Added use_hipporag to arguments
    def __init__(self, embedding_model_name: str = None, use_hipporag: bool = False):
        """
        Inizializza il mapper migliorato.
        
        Args:
            embedding_model_name: Nome del modello per embeddings
            use_hipporag: Abilita il motore associativo (Graph)
        """
        # Load model names from environment variables
        if embedding_model_name is None:
            embedding_model_name = os.getenv("EMBEDDING_MODEL", "paraphrase-multilingual-MiniLM-L12-v2")
        print(f"🔄 Inizializzazione sistema di mapping migliorato...")
        
        # Modello di embedding
        self.embedding_model = SentenceTransformer(embedding_model_name)
        print(f"✅ Modello embedding caricato: {embedding_model_name}")
        
        # TF-IDF per keyword matching
        self.tfidf_vectorizer = None
        self.tfidf_matrix = None
        
        # Cache per ottimizzazione
        self.embedding_cache = {}
        
        # Stopwords italiane
        self.stopwords = set([
            'il', 'lo', 'la', 'i', 'gli', 'le', 'un', 'uno', 'una', 'di', 'da', 'in', 'su',
            'per', 'con', 'tra', 'fra', 'a', 'e', 'è', 'o', 'ma', 'se', 'che', 'chi', 'cui'
        ])

        # --- [FIXED] HippoRAG Setup ---
        self.use_hipporag = use_hipporag
        self.hippo_engine = None 

        if self.use_hipporag:
            print("🧠 Abilitazione HippoRAG Lite (Associative Memory)...")
            try:
                # Initialize the Lite engine
                self.hippo_engine = HippoRAGLite()
                print("✅ HippoRAG Lite inizializzato.")
            except Exception as e:
                print(f"⚠️ Errore inizializzazione HippoRAG Lite: {e}")
                self.use_hipporag = False

    # [FIX 2] Fixed Indentation (Inside Class)
    def build_hipporag_index(self, chunks: List[Dict[str, Any]]):
        """
        Costruisce il grafo associativo usando il motore Lite.
        """
        if not self.use_hipporag or not self.hippo_engine:
            return
            
        print("\n🧠 Costruzione Indice HippoRAG Lite...")
        try:
            # The Lite engine builds the graph in memory immediately (fast)
            self.hippo_engine.build_index(chunks)
        except Exception as e:
            print(f"❌ Errore durante l'indexing HippoRAG: {e}")
            self.use_hipporag = False

    def load_requirements_enhanced(self, requirements_file: str, max_requirements: int = None, 
                                  start_from: int = 0) -> List[Dict[str, Any]]:
        """
        Carica requisiti con parsing migliorato e limitazione opzionale.
        """
        requirements = []
        
        if not os.path.exists(requirements_file):
            print(f"❌ File requisiti non trovato: {requirements_file}")
            return requirements
        
        try:
            with open(requirements_file, 'r', encoding='utf-8') as f:
                content = f.read()
            
            req_blocks = re.split(r'REQUISITO \d+', content)[1:]
            
            # Applica limitazioni se specificate
            if max_requirements or start_from > 0:
                end_idx = start_from + max_requirements if max_requirements else len(req_blocks)
                req_blocks = req_blocks[start_from:end_idx]
                print(f"⚠️  Caricamento limitato: requisiti {start_from+1} a {min(end_idx, len(req_blocks)+start_from)}")
            
            for i, block in enumerate(req_blocks, start_from + 1):
                lines = block.strip().split('\n')
                
                requirement_data = {
                    'id': i,
                    'requirement_original': "",
                    'requirement_paraphrased': "",
                    'audit_question': "",
                    'source_file': "",
                    'source_chunk': "",
                    'keywords': [],
                    'requirement_type': "",
                    'references': ""
                }
                
                for line in lines:
                    line = line.strip()
                    if line.startswith('File sorgente:'):
                        requirement_data['source_file'] = line.replace('File sorgente:', '').strip()
                    elif line.startswith('Chunk sorgente:'):
                        requirement_data['source_chunk'] = line.replace('Chunk sorgente:', '').strip()
                    elif line.startswith('REQUISITO ORIGINALE:'):
                        requirement_data['requirement_original'] = line.replace('REQUISITO ORIGINALE:', '').strip()
                    elif line.startswith('REQUISITO PARAFRASATO:'):
                        requirement_data['requirement_paraphrased'] = line.replace('REQUISITO PARAFRASATO:', '').strip()
                    elif line.startswith('REQUISITO:'):
                        requirement_data['requirement_paraphrased'] = line.replace('REQUISITO:', '').strip()
                    elif line.startswith('DOMANDA AUDIT:'):
                        requirement_data['audit_question'] = line.replace('DOMANDA AUDIT:', '').strip()
                    elif line.startswith('REFERENZE:'):
                        requirement_data['references'] = line.replace('REFERENZE:', '').strip()
                
                # Fallback
                if not requirement_data['requirement_original'] and requirement_data['requirement_paraphrased']:
                    requirement_data['requirement_original'] = requirement_data['requirement_paraphrased']
                
                # Estrai keywords e tipo di requisito
                if requirement_data['requirement_original']:
                    requirement_data['keywords'] = self.extract_keywords(requirement_data['requirement_original'])
                    requirement_data['requirement_type'] = self.classify_requirement_type(requirement_data['requirement_original'])
                    requirement_data['requirement_combined'] = f"{requirement_data['requirement_original']} {requirement_data['requirement_paraphrased']}".strip()
                    requirements.append(requirement_data)
            
            print(f"📋 Caricati {len(requirements)} requisiti con metadati arricchiti")
            return requirements
            
        except Exception as e:
            print(f"❌ Errore nel caricare i requisiti: {e}")
            return []
    
    def load_improved_chunks(self, chunks_dir: str) -> List[Dict[str, Any]]:
        """
        Carica chunks dal formato migliorato (JSON).
        """
        all_chunks = []
        
        chunk_files = list(Path(chunks_dir).glob("*_improved_chunks.json"))
        
        if not chunk_files:
            # Prova formato alternativo
            chunk_files = list(Path(chunks_dir).glob("*.json"))
        
        for chunk_file in chunk_files:
            print(f"📁 Caricamento chunks da: {chunk_file.name}")
            
            try:
                with open(chunk_file, 'r', encoding='utf-8') as f:
                    chunks = json.load(f)
                
                # Aggiungi indice globale
                for chunk in chunks:
                    chunk['global_id'] = len(all_chunks)
                    # Estrai numero sezione se presente
                    chunk['section_number'] = self.extract_section_number(chunk.get('section', ''))
                    all_chunks.append(chunk)
                
                print(f"   ✅ Caricati {len(chunks)} chunks")
                
            except Exception as e:
                print(f"   ❌ Errore: {e}")
                continue
        
        print(f"📊 Totale chunks caricati: {len(all_chunks)}")
        return all_chunks
    
    def extract_section_number(self, section_title: str) -> str:
        """
        Estrae il numero di sezione dal titolo.
        """
        # Pattern per numeri di sezione comuni
        patterns = [
            r'^(\d+\.?\d*\.?\d*)',  # 1.2.3 o 1.2 o 1
            r'^([IVXLCDM]+)',       # Numeri romani
            r'^([A-Z])\.',          # A. B. C.
        ]
        
        for pattern in patterns:
            match = re.match(pattern, section_title)
            if match:
                return match.group(1)
        
        return ""
    
    def extract_keywords(self, text: str) -> List[str]:
        """
        Estrae keywords significative dal testo.
        """
        # Tokenizzazione base
        words = re.findall(r'\b[a-zA-ZàèéìòùÀÈÉÌÒÙ]{3,}\b', text.lower())
        
        # Rimuovi stopwords
        keywords = [w for w in words if w not in self.stopwords]
        
        # Identifica termini importanti
        important_patterns = [
            r'deve\s+(\w+)',
            r'necessario\s+(\w+)',
            r'richiesto\s+(\w+)',
            r'obbligo\s+di\s+(\w+)',
            r'divieto\s+di\s+(\w+)'
        ]
        
        for pattern in important_patterns:
            matches = re.findall(pattern, text.lower())
            keywords.extend(matches)
        
        return list(set(keywords))[:20]
    
    def classify_requirement_type(self, text: str) -> str:
        """
        Classifica il tipo di requisito.
        """
        text_lower = text.lower()
        
        if any(word in text_lower for word in ['documentare', 'registrare', 'policy', 'procedura']):
            return 'documentazione'
        elif any(word in text_lower for word in ['formare', 'formazione', 'training', 'addestramento']):
            return 'formazione'
        elif any(word in text_lower for word in ['monitorare', 'verificare', 'audit', 'controllo']):
            return 'monitoraggio'
        elif any(word in text_lower for word in ['comunicare', 'informare', 'notificare']):
            return 'comunicazione'
        elif any(word in text_lower for word in ['non deve', 'vietato', 'proibito']):
            return 'divieto'
        else:
            return 'generale'
    
    def compute_tfidf_scores(self, query: str, documents: List[str]) -> np.ndarray:
        """
        Calcola score TF-IDF tra query e documenti.
        """
        if not self.tfidf_vectorizer:
            self.tfidf_vectorizer = TfidfVectorizer(
                max_features=1000,
                ngram_range=(1, 2),
                stop_words=list(self.stopwords)
            )
            
            # Fit su tutti i documenti
            all_texts = [query] + documents
            self.tfidf_matrix = self.tfidf_vectorizer.fit_transform(all_texts)
        else:
            # Transform query
            query_vec = self.tfidf_vectorizer.transform([query])
            docs_vecs = self.tfidf_vectorizer.transform(documents)
            
            # Calcola similarità
            scores = cosine_similarity(query_vec, docs_vecs)[0]
            return scores
        
        # Prima volta: estrai scores
        query_vec = self.tfidf_matrix[0]
        docs_vecs = self.tfidf_matrix[1:]
        scores = cosine_similarity(query_vec, docs_vecs)[0]
        
        return scores
    
    def keyword_overlap_score(self, req_keywords: List[str], chunk_text: str) -> float:
        """
        Calcola score basato su overlap di keywords.
        """
        if not req_keywords:
            return 0.0
        
        chunk_words = set(re.findall(r'\b[a-zA-ZàèéìòùÀÈÉÌÒÙ]{3,}\b', chunk_text.lower()))
        req_keywords_set = set(req_keywords)
        
        overlap = len(req_keywords_set.intersection(chunk_words))
        
        # Normalizza per numero di keywords
        score = overlap / len(req_keywords_set)
        
        return score

    # [FIX 2] Fixed Indentation (Inside Class)
    def hipporag_retrieval(self, requirement: Dict[str, Any], top_k: int = 15) -> List[Dict[str, Any]]:
        """
        Esegue il retrieval associativo usando HippoRAG Lite.
        """
        # [FIX 4] Added safety check
        if not self.use_hipporag or not self.hippo_engine:
            return []
            
        # Combine original and paraphrased text for better entity extraction
        query = f"{requirement['requirement_original']} {requirement['requirement_paraphrased']}"
        
        try:
            # Call the Lite engine directly
            results = self.hippo_engine.retrieve(query, top_k=top_k)
            
            matches = []
            for i, res in enumerate(results):
                chunk = res['chunk']
                score = res['score']
                entities = res.get('matched_entities', [])
                
                matches.append({
                    'chunk': chunk,
                    'similarity_score': float(score), # Use graph score as similarity
                    'tfidf_score': 0.0,
                    'keyword_score': 0.0,
                    'final_score': float(score),
                    'rank': i + 1,
                    'source': f"hipporag (via {', '.join(entities[:3])})" # Traceability
                })
                
            if matches:
                print(f"      🧠 HippoRAG attivato: {len(matches)} chunks collegati.")
                
            return matches
            
        except Exception as e:
            print(f"      ⚠️ Errore durante HippoRAG retrieval: {e}")
            return []
    
    def hybrid_retrieval(self, requirement: Dict[str, Any], chunks: List[Dict[str, Any]], 
                        req_embedding: np.ndarray, chunk_embeddings: np.ndarray,
                        top_k: int = 30) -> List[Dict[str, Any]]:
        """
        Retrieval ibrido che combina diversi metodi di scoring.
        """
        print(f"   🔍 Hybrid retrieval per requisito {requirement['id']}...")
        
        # 1. Semantic similarity scores
        semantic_scores = cosine_similarity([req_embedding], chunk_embeddings)[0]
        
        # 2. TF-IDF scores
        chunk_texts = [c.get('full_text_with_context', c['text']) for c in chunks]
        tfidf_scores = self.compute_tfidf_scores(requirement['requirement_combined'], chunk_texts)
        
        # 3. Keyword overlap scores
        keyword_scores = np.array([
            self.keyword_overlap_score(requirement['keywords'], c.get('full_text_with_context', c['text']))
            for c in chunks
        ])
        
        # 4. Boost per sezioni rilevanti
        section_boost = np.array([
            1.2 if self.is_relevant_section(requirement['requirement_type'], c.get('section', ''))
            else 1.0
            for c in chunks
        ])

        # [NEW] Get Hippo Scores with Safety Check
        hippo_score_vector = np.zeros(len(chunks))

        if self.use_hipporag and self.hippo_engine:
            hippo_matches = self.hipporag_retrieval(requirement)
            
            for res in hippo_matches:
                c_id = res['chunk']['global_id']
                # Safe mapping check
                if c_id < len(chunks):
                    hippo_score_vector[c_id] = res['final_score']

        # Combina scores con pesi
        final_scores = (
            0.40 * semantic_scores + 
            0.20 * tfidf_scores + 
            0.10 * keyword_scores + 
            0.30 * hippo_score_vector # Add associative score
        ) * section_boost
        
        # Normalizza scores
        if final_scores.max() > 0:
            final_scores = final_scores / final_scores.max()
        
        # Seleziona top K
        top_indices = np.argsort(final_scores)[::-1][:top_k]
        
        matches = []
        for idx in top_indices:
            if final_scores[idx] > 0.25:
                # Determine source label for clarity
                src = "hybrid"
                if hippo_score_vector[idx] > 0:
                    src += "+hipporag"

                matches.append({
                    'chunk': chunks[idx],
                    'similarity_score': float(semantic_scores[idx]),
                    'tfidf_score': float(tfidf_scores[idx]),
                    'keyword_score': float(keyword_scores[idx]),
                    'final_score': float(final_scores[idx]),
                    'rank': len(matches) + 1,
                    'source': src
                })
        
        print(f"      ✅ Trovati {len(matches)} matches con score > 0.25")
        
        # Re-ranking finale
        matches = sorted(matches, key=lambda x: x['final_score'], reverse=True)[:15]
        
        return matches
    
    def is_relevant_section(self, req_type: str, section_title: str) -> bool:
        """
        Verifica se una sezione è rilevante per il tipo di requisito.
        """
        section_lower = section_title.lower()
        
        relevance_map = {
            'documentazione': ['documento', 'procedura', 'policy', 'registro'],
            'formazione': ['formazione', 'training', 'addestramento', 'competenze'],
            'monitoraggio': ['monitoraggio', 'audit', 'verifica', 'controllo'],
            'comunicazione': ['comunicazione', 'informazione', 'notifica'],
            'divieto': ['divieto', 'restrizione', 'limitazione']
        }
        
        relevant_terms = relevance_map.get(req_type, [])
        return any(term in section_lower for term in relevant_terms)
    
    def enhanced_llm_analysis(self, requirement: Dict[str, Any], matches: List[Dict[str, Any]],
                             model_name: str = None) -> Dict[str, Any]:
        """
        Analisi LLM migliorata con valutazione strutturata.
        """
        # Get model name from environment if not provided
        if model_name is None:
            model_name = os.getenv("LLM_MAPPING_MODEL", "deepseek-r1:14b")
        
        if not matches:
            return {
                "allineamento": "Nessun chunk rilevante trovato nei documenti interni.",
                "gap": "Requisito completamente non coperto.",
                "livello_copertura": 1,
                "descrizione_livello": "ASSENTE",
                "confidence": 1.0,
                "aspetti_coperti": [],
                "aspetti_mancanti": ["Tutti gli aspetti del requisito"],
                "azioni_richieste": []
            }
        
        # Prepara contesto arricchito
        chunks_context = self.prepare_enriched_context(matches)
        
        # Usa i prompt dal file di configurazione
        system_prompt = MAPPING_SYSTEM_PROMPT
        user_prompt = get_mapping_user_prompt(
            requirement_original=requirement['requirement_original'],
            requirement_type=requirement.get('requirement_type', 'generale'),
            keywords=requirement.get('keywords', []),
            chunks_context=chunks_context
        )
        
        try:
            print(
            f"DEBUG LLM CALL → provider={os.getenv('LLM_PROVIDER')} "
            f"model={model_name}"
            )

            # Use llm_provider instead of direct ollama call
            response = llm_chat(
                model=model_name,
                messages=[
                    {'role': 'system', 'content': system_prompt},
                    {'role': 'user', 'content': user_prompt}
                ],
                options={
                    "num_predict": int(os.getenv("LLM_NUM_PREDICT", "512")),
                    "temperature": float(os.getenv("LLM_MAPPING_TEMPERATURE", "0.1")),
                    "num_ctx": int(os.getenv("LLM_NUM_CTX", "8192")),
                    "top_p": float(os.getenv("LLM_TOP_P", "0.95"))
                }
            )
            
            response_text = response.get("message", {}).get("content", "").strip()
            
            # Parse JSON
            json_match = re.search(r'\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}', response_text, re.DOTALL)
            if json_match:
                analysis = json.loads(json_match.group(0))
                
                # Validazione e normalizzazione
                analysis = self.validate_llm_response(analysis)
                return analysis
            else:
                raise ValueError("No valid JSON in response")
                
        except Exception as e:
            print(f"      ❌ Errore analisi LLM: {e}")
            return self.get_default_analysis()
    
    def prepare_enriched_context(self, matches: List[Dict[str, Any]]) -> str:
        """
        Prepara contesto arricchito per l'analisi LLM.
        """
        context_parts = []
        
        for i, match in enumerate(matches[:10], 1):
            chunk = match['chunk']
            
            # Prepara informazioni sulla sezione
            section_title = chunk.get('section', 'N/A')
            section_number = chunk.get('section_number', '')
            
            # Se il numero di sezione è già nel titolo, non duplicarlo
            if section_number and section_title.startswith(section_number):
                section_info = section_title
            elif section_number:
                section_info = f"{section_number}. {section_title}"
            else:
                section_info = section_title
            
            # Add source label
            source_label = match.get('source', 'hybrid')

            context_parts.append(f"""
DOCUMENTO {i} (Source: {source_label} - Score: {match['final_score']:.2f})
File: {chunk.get('source_file', 'N/A')}
Sezione: {section_info}
Livello gerarchico: {chunk.get('hierarchy_level', 0)}

CONTENUTO:
{chunk.get('full_text_with_context', chunk['text'])}

---""")
        
        return '\n'.join(context_parts)
    
    def validate_llm_response(self, analysis: Dict[str, Any]) -> Dict[str, Any]:
        """
        Valida e normalizza la risposta LLM.
        """
        # Campi richiesti
        required = ["allineamento", "gap", "livello_copertura", "descrizione_livello"]
        
        for field in required:
            if field not in analysis:
                analysis[field] = self.get_default_value(field)
        
        # Valida livello copertura
        if not isinstance(analysis.get("livello_copertura"), int):
            analysis["livello_copertura"] = 1
        
        analysis["livello_copertura"] = max(1, min(5, analysis["livello_copertura"]))
        
        # Normalizza descrizione livello
        level_map = {
            5: "COMPLETA",
            4: "ELEVATA", 
            3: "MEDIA",
            2: "BASSA",
            1: "ASSENTE"
        }
        
        analysis["descrizione_livello"] = level_map.get(analysis["livello_copertura"], "ASSENTE")
        
        # Aggiungi campi opzionali se mancanti
        analysis.setdefault("confidence", 0.8)
        analysis.setdefault("aspetti_coperti", [])
        analysis.setdefault("aspetti_mancanti", [])
        analysis.setdefault("azioni_richieste", [])
        
        return analysis
    
    def get_default_value(self, field: str) -> Any:
        """
        Ritorna valore di default per campo.
        """
        defaults = {
            "allineamento": "Analisi non disponibile",
            "gap": "Gap non identificati",
            "livello_copertura": 1,
            "descrizione_livello": "ASSENTE",
            "confidence": 0.5
        }
        return defaults.get(field, "")
    
    def get_default_analysis(self) -> Dict[str, Any]:
        """
        Ritorna analisi di default in caso di errore.
        """
        return {
            "allineamento": "Errore nell'analisi automatica",
            "gap": "Impossibile determinare gap",
            "livello_copertura": 1,
            "descrizione_livello": "ASSENTE",
            "confidence": 0.0,
            "aspetti_coperti": [],
            "aspetti_mancanti": ["Analisi non disponibile"],
            "azioni_richieste": ["Verificare manualmente"]
        }
    
    def clean_text_artifacts(self, text: str) -> str:
        """
        Pulisce artefatti di estrazione PDF tipo 'sliding window'.
        """
        if not text or len(text) < 3:
            return text
            
        tokens = text.split()
        if not tokens:
            return ""
            
        output = []
        current_word = tokens[0]
        
        for i in range(1, len(tokens)):
            token = tokens[i]
            if current_word and token and current_word[-1] == token[0]:
                current_word += token[1:]
            else:
                output.append(current_word)
                current_word = token
                
        output.append(current_word)
        return " ".join(output)

    def format_top_matches_for_excel(self, matches: List[Dict[str, Any]]) -> str:
        """
        Formatta i top matches per la visualizzazione in Excel.
        """
        if not matches:
            return "Nessun match trovato"
        
        formatted_matches = []
        for i, match in enumerate(matches[:10], 1):
            chunk = match['chunk']
            
            section_title = chunk.get('section', 'N/A')[:50]
            section_number = chunk.get('section_number', '')
            
            if section_number and section_title.startswith(section_number):
                section_info = section_title
            elif section_number:
                section_info = f"{section_number}. {section_title}"
            else:
                section_info = section_title
            
            section_info = self.clean_text_artifacts(section_info)
            
            match_text = (
                f"{i}. File: {chunk.get('source_file', 'N/A')}\n"
                f"   Pagina: {chunk.get('page_number', 'N/A')}\n"
                f"   Sezione: {section_info}\n"
                f"   Score: {match['final_score']:.3f}"
            )
            formatted_matches.append(match_text)
        
        return "\n\n".join(formatted_matches)
    
    def create_excel_output(self, results: List[Dict[str, Any]], output_path: str):
        """
        Crea output Excel con tutte le colonne richieste.
        """
        # Prepara dati per DataFrame
        excel_data = []
        
        for result in results:
            analysis = result['analysis']
            
            row = {
                'REQUISITO ESTERNO': result['requirement_original'],
                'REFERENZE': result.get('references', ''),
                'DOMANDA AUDIT': result.get('audit_question', ''),
                'MAPPING DOCUMENTI INTERNI': self.format_top_matches_for_excel(result['matches']),
                'LIVELLO': f"{analysis['livello_copertura']} - {analysis['descrizione_livello']}",
                'ALLINEAMENTO': analysis['allineamento'],
                'GAP': analysis['gap'],
                'ASPETTI COPERTI': ', '.join(analysis.get('aspetti_coperti', [])),
                'ASPETTI MANCANTI': ', '.join(analysis.get('aspetti_mancanti', [])),
                'AZIONI RICHIESTE': ', '.join(analysis.get('azioni_richieste', []))
            }
            excel_data.append(row)
        
        df = pd.DataFrame(excel_data)
        excel_file = Path(output_path) / "mapping_results.xlsx"
        
        with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Risultati Mapping', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Risultati Mapping']
            
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            column_widths = {
                'A': 50, 'B': 30, 'C': 40, 'D': 60, 'E': 20, 
                'F': 50, 'G': 50, 'H': 40, 'I': 40, 'J': 40
            }
            
            for column, width in column_widths.items():
                worksheet.column_dimensions[column].width = width
            
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            data_alignment = Alignment(vertical="top", wrap_text=True)
            
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = data_alignment
                    cell.border = thin_border
                
                level_cell = row[4]
                if level_cell.value:
                    if '5 -' in str(level_cell.value):
                        level_cell.fill = PatternFill(start_color="0EC000", end_color="0EC000", fill_type="solid")
                    elif '4 -' in str(level_cell.value):
                        level_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    elif '3 -' in str(level_cell.value):
                        level_cell.fill = PatternFill(start_color="FAE050", end_color="FAE050", fill_type="solid")
                    elif '2 -' in str(level_cell.value):
                        level_cell.fill = PatternFill(start_color="EA9050", end_color="EA9050", fill_type="solid")
                    elif '1 -' in str(level_cell.value):
                        level_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                        level_cell.font = Font(color="FFFFFF")
            
            worksheet.auto_filter.ref = worksheet.dimensions
        
        print(f"✅ File Excel salvato: {excel_file}")
    
    def process_enhanced_mapping(self, requirements_file: str, chunks_dir: str,
                                output_dir: str, use_combined_embeddings: bool = True,
                                max_requirements: int = None, start_from: int = 0):
        """
        Pipeline completa di mapping migliorato.
        """
        print("\n" + "="*80)
        print("🚀 AVVIO MAPPING COMPLIANCE MIGLIORATO")
        print(f"   Configurazione: {max_requirements if max_requirements else 'tutti i'} requisiti dal #{start_from+1}")
        print("="*80)
        
        Path(output_dir).mkdir(exist_ok=True)
        
        # 1. Carica requisiti
        print("\n📋 FASE 1: Caricamento requisiti...")
        requirements = self.load_requirements_enhanced(requirements_file, max_requirements, start_from)
        if not requirements: return {"status": "error", "message": "Nessun requisito caricato"}
        
        # 2. Carica chunks
        print("\n📁 FASE 2: Caricamento chunks...")
        chunks = self.load_improved_chunks(chunks_dir)
        if not chunks: return {"status": "error", "message": "Nessun chunk caricato"}
        
        # 3. Prepara TF-IDF
        print("\n📊 FASE 3: Preparazione TF-IDF...")
        all_texts = [r['requirement_combined'] for r in requirements]
        all_texts.extend([c.get('full_text_with_context', c['text']) for c in chunks])
        
        self.tfidf_vectorizer = TfidfVectorizer(max_features=1000, ngram_range=(1, 2), stop_words=list(self.stopwords))
        self.tfidf_matrix = self.tfidf_vectorizer.fit_transform(all_texts)
        
        # [FIX 3] Call Build Index Here
        if self.use_hipporag:
            self.build_hipporag_index(chunks)

        # 4. Calcola embeddings
        print("\n🧮 FASE 4: Calcolo embeddings...")
        
        if use_combined_embeddings:
            print("   Usando embeddings combinati (originale + parafrasi)...")
            req_embeddings = []
            for req in requirements:
                emb_orig = self.embedding_model.encode(req['requirement_original'])
                emb_para = self.embedding_model.encode(req['requirement_paraphrased'])
                
                if req['requirement_original'] != req['requirement_paraphrased']:
                    combined = 0.6 * emb_orig + 0.4 * emb_para
                else:
                    combined = emb_orig
                req_embeddings.append(combined)
            req_embeddings = np.array(req_embeddings)
        else:
            req_texts = [r['requirement_combined'] for r in requirements]
            req_embeddings = self.embedding_model.encode(req_texts, show_progress_bar=True)
        
        chunk_texts = [c.get('full_text_with_context', c['text']) for c in chunks]
        chunk_embeddings = self.embedding_model.encode(chunk_texts, show_progress_bar=True)
        
        # 5. Process mapping
        print(f"\n🔄 FASE 5: Processamento {len(requirements)} requisiti...")
        results = []
        
        for i, requirement in enumerate(requirements):
            print(f"\n📍 Requisito {i+1}/{len(requirements)} (ID: {requirement['id']})")
            print(f"   {requirement['requirement_original'][:100]}...")
            
            # Hybrid retrieval
            matches = self.hybrid_retrieval(
                requirement, 
                chunks,
                req_embeddings[i],
                chunk_embeddings,
                top_k=30
            )
            
            # Analisi LLM
            if matches:
                print("   🤖 Analisi LLM in corso...")
                analysis = self.enhanced_llm_analysis(requirement, matches)
                print(f"   📊 Risultato: Livello {analysis['livello_copertura']} - {analysis['descrizione_livello']}")
            else:
                analysis = self.get_default_analysis()
            
            result = {
                "requirement_id": requirement['id'],
                "requirement_original": requirement['requirement_original'],
                "requirement_paraphrased": requirement.get('requirement_paraphrased', ''),
                "requirement_type": requirement.get('requirement_type', ''),
                "keywords": requirement.get('keywords', []),
                "audit_question": requirement.get('audit_question', ''),
                "references": requirement.get('references', ''),
                "source_info": {
                    "file": requirement.get('source_file', ''),
                    "chunk": requirement.get('source_chunk', '')
                },
                "matches": matches,
                "analysis": analysis
            }
            results.append(result)
        
        # 6. Salva risultati
        print(f"\n💾 FASE 6: Salvataggio risultati...")
        self.save_enhanced_results(results, output_dir)
        print("📊 Creazione file Excel...")
        self.create_excel_output(results, output_dir)
        
        # 7. Calcola statistiche
        stats = self.calculate_enhanced_statistics(results)
        
        print(f"\n" + "="*80)
        print("📊 STATISTICHE FINALI:")
        print("="*80)
        print(f"Requisiti processati: {stats['total_requirements']}")
        print(f"Copertura media: {stats['average_coverage']:.1f}%")
        print(f"\nDistribuzione livelli:")
        for level in range(5, 0, -1):
            count = stats['coverage_levels'].get(f'level_{level}', 0)
            perc = stats['coverage_percentages'].get(f'level_{level}', '0%')
            print(f"   Livello {level}: {count} requisiti ({perc})")
        
        print(f"\nTipi di requisiti:")
        for req_type, count in stats.get('requirement_types', {}).items():
            print(f"   {req_type}: {count}")
        
        return {
            "status": "success",
            "stats": stats,
            "output_dir": output_dir
        }
    
    def save_enhanced_results(self, results: List[Dict[str, Any]], output_dir: str):
        """
        Salva risultati in formato JSON e testo.
        """
        output_path = Path(output_dir)
        
        json_file = output_path / "mapping_results.json"
        with open(json_file, 'w', encoding='utf-8') as f:
            json.dump(results, f, ensure_ascii=False, indent=2)
        
        txt_file = output_path / "mapping_report.txt"
        with open(txt_file, 'w', encoding='utf-8') as f:
            f.write("REPORT MAPPING COMPLIANCE - ANALISI MIGLIORATA\n")
            f.write("="*80 + "\n\n")
            
            for result in results:
                f.write(f"REQUISITO {result['requirement_id']}\n")
                f.write("-"*60 + "\n")
                f.write(f"TESTO: {result['requirement_original']}\n")
                f.write(f"TIPO: {result.get('requirement_type', 'N/A')}\n")
                if result.get('audit_question'):
                    f.write(f"DOMANDA AUDIT: {result['audit_question']}\n")
                f.write("\n")
                
                f.write("TOP MATCHES:\n")
                for match in result['matches'][:5]:
                    chunk = match['chunk']
                    section_title = chunk.get('section', 'N/A')
                    section_number = chunk.get('section_number', '')
                    
                    if section_number and section_title.startswith(section_number):
                        section_info = section_title
                    elif section_number:
                        section_info = f"{section_number}. {section_title}"
                    else:
                        section_info = section_title
                    
                    f.write(f"   {match['rank']}. Score: {match['final_score']:.3f}\n")
                    f.write(f"     File: {chunk.get('source_file', 'N/A')}\n")
                    f.write(f"     Sezione: {section_info}\n")
                
                f.write("\nANALISI:\n")
                analysis = result['analysis']
                f.write(f"   Livello: {analysis['livello_copertura']} - {analysis['descrizione_livello']}\n")
                f.write(f"   Allineamento: {analysis['allineamento']}\n")
                f.write(f"   Gap: {analysis['gap']}\n")
                
                if analysis.get('aspetti_coperti'):
                    f.write(f"   Aspetti coperti: {', '.join(analysis['aspetti_coperti'])}\n")
                if analysis.get('aspetti_mancanti'):
                    f.write(f"   Aspetti mancanti: {', '.join(analysis['aspetti_mancanti'])}\n")
                if analysis.get('azioni_richieste'):
                    f.write(f"   Azioni richieste: {', '.join(analysis['azioni_richieste'])}\n")
                
                f.write("\n" + "="*80 + "\n\n")
        
        print(f"✅ Risultati salvati in {output_dir}/")
    
    def calculate_enhanced_statistics(self, results: List[Dict[str, Any]]) -> Dict[str, Any]:
        """
        Calcola statistiche dettagliate.
        """
        total = len(results)
        
        if total == 0:
            return {"total_requirements": 0, "average_coverage": 0}
        
        level_counts = {}
        for level in range(1, 6):
            count = sum(1 for r in results if r['analysis']['livello_copertura'] == level)
            level_counts[f'level_{level}'] = count
        
        level_percentages = {}
        for level, count in level_counts.items():
            level_percentages[level] = f"{(count/total*100):.1f}%"
        
        coverage_scores = []
        for r in results:
            perc = r['analysis'].get('percentuale_copertura', 0)
            if isinstance(perc, (int, float)):
                coverage_scores.append(perc)
            else:
                level = r['analysis']['livello_copertura']
                estimated = {1: 10, 2: 35, 3: 60, 4: 80, 5: 95}
                coverage_scores.append(estimated.get(level, 0))
        
        avg_coverage = np.mean(coverage_scores) if coverage_scores else 0
        
        req_types = {}
        for r in results:
            req_type = r.get('requirement_type', 'generale')
            req_types[req_type] = req_types.get(req_type, 0) + 1
        
        return {
            "total_requirements": total,
            "coverage_levels": level_counts,
            "coverage_percentages": level_percentages,
            "average_coverage": avg_coverage,
            "requirement_types": req_types,
            "critical_requirements": sum(1 for r in results if r['analysis']['livello_copertura'] <= 2),
            "well_covered": sum(1 for r in results if r['analysis']['livello_copertura'] >= 4)
        }

def main():
    """
    Pipeline completa di esecuzione.
    """
    # Configurazione
    mapper = ImprovedRequirementsMapper(
        embedding_model_name="paraphrase-multilingual-MiniLM-L12-v2",
        use_hipporag=True  # [FIX 5] Enable HippoRAG here
    )
    
    # Percorsi
    requirements_file = "/home/ashesh_kumar_gupta/output_1712/final_requirements/req_2.05.23_FAQ_Accredia_PdR125_final.txt"
    chunks_dir = "/home/ashesh_kumar_gupta/internal"
    output_dir = "/home/ashesh_kumar_gupta/output/mappingFAQ_Accredia_PdR125_improved_hippo"
    
    # CONFIGURAZIONE: numero requisiti da processare
    MAX_REQUIREMENTS = None 
    START_FROM = 0
    
    print(f"\n⚙️  CONFIGURAZIONE:")
    print(f"   - Requisiti da processare: {'TUTTI' if not MAX_REQUIREMENTS else MAX_REQUIREMENTS}")
    print(f"   - Inizia dal requisito: {START_FROM + 1}")
    
    # Esegui mapping
    result = mapper.process_enhanced_mapping(
        requirements_file=requirements_file,
        chunks_dir=chunks_dir,
        output_dir=output_dir,
        use_combined_embeddings=True,
        max_requirements=MAX_REQUIREMENTS,
        start_from=START_FROM
    )
    
    if result["status"] == "success":
        print(f"\n✅ PROCESSO COMPLETATO CON SUCCESSO!")
        print(f"📁 Risultati salvati in: {result['output_dir']}")
        print(f"📊 File Excel generato: mapping_results.xlsx")
    else:
        print(f"\n❌ Processo fallito: {result.get('message', 'Errore sconosciuto')}")

if __name__ == "__main__":
    main()